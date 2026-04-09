"""
Phase 2: 文件内容提取
按后缀分派不同解析器，统一返回纯文本
"""
import chardet
from pathlib import Path


MAX_CHARS = 8000  # 单文件最多读取字符数


def read_file(path: str) -> str | None:
    """读取文件内容，返回纯文本，失败返回 None"""
    ext = Path(path).suffix.lower()
    try:
        if ext in (".md", ".txt", ".csv", ".sql", ".r", ".m"):
            return _read_text(path)
        elif ext in (".py", ".ipynb"):
            return _read_text(path)
        elif ext in (".docx", ".doc"):
            return _read_docx(path)
        elif ext == ".pdf":
            return _read_pdf(path)
        elif ext in (".xlsx", ".xls"):
            return _read_excel(path)
        elif ext == ".json":
            return _read_json_as_text(path)
    except Exception as e:
        print(f"[FileReader] 读取失败 {path}: {e}")
    return None


def _read_text(path: str) -> str:
    raw = open(path, "rb").read()
    enc = chardet.detect(raw).get("encoding") or "utf-8"
    text = raw.decode(enc, errors="ignore")
    return text[:MAX_CHARS]


def _read_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(lines)[:MAX_CHARS]


def _read_pdf(path: str) -> str:
    import pdfplumber
    texts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages[:10]:   # 最多读10页
            t = page.extract_text()
            if t:
                texts.append(t)
            if sum(len(t) for t in texts) > MAX_CHARS:
                break
    return "\n".join(texts)[:MAX_CHARS]


def _read_excel(path: str) -> str:
    """只读取前几行，了解表格结构和内容"""
    import csv, io
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets[:2]:
            lines.append(f"[Sheet: {sheet.title}]")
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i > 20:
                    break
                lines.append("\t".join(str(c) for c in row if c is not None))
        return "\n".join(lines)[:MAX_CHARS]
    except ImportError:
        return None


def _read_json_as_text(path: str) -> str:
    """只读小 JSON 文件，且跳过明显是配置/依赖的"""
    import json
    skip_names = {"package.json", "package-lock.json", "tsconfig.json",
                  "composer.json", ".eslintrc.json", "launch.json"}
    if Path(path).name in skip_names:
        return None
    raw = open(path, "rb").read()
    if len(raw) > 50 * 1024:   # 跳过 > 50KB 的 json
        return None
    enc = chardet.detect(raw).get("encoding") or "utf-8"
    return raw.decode(enc, errors="ignore")[:MAX_CHARS]
